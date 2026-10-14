#!/usr/bin/env python3
"""Tira da ficha do Ciclo Maldito o que o bot precisa da aba ANALISE. A ficha é a dona do contrato: nada aqui é escrito à
mão, e se a ficha mudar é só rodar de novo.

    python ferramentas/ficha_analise.py chaves FICHA.xlsx|LINK   # gera dados/ficha/analise-chaves.tsv
    python ferramentas/ficha_analise.py amostra FICHA.xlsx [SAIDA] ["ABA!CÉLULA=valor" ...]
                                                                  # preenche, recalcula no LibreOffice e grava o caso de teste

`chaves` lê a coluna A (chave), a C (célula) e a D (parte) da aba, do .xlsx gerado (sem precisar de valores) ou do
link aberto. `amostra` precisa dos valores: copia a ficha, recalcula no LibreOffice e grava só a aba ANALISE, com os
valores, em src/testes/dados/ (padrão: ficha_analise_branca.xlsx). Com "FICHA!D11=Bastião" e afins, escreve as
entradas antes de recalcular: é assim que saem as fichas de teste preenchidas, calculadas pela própria ficha.
O LibreOffice pesa: rode um por vez e com teto, assim:

    systemd-run --user --scope -q -p MemoryMax=5G -p MemorySwapMax=0 nice -n 15 \\
        .venv/bin/python ferramentas/ficha_analise.py amostra FICHA.xlsx
"""
from __future__ import annotations

import asyncio
import os
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import date
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "src"))

from closure_bot.ficha_leitor import ABA, CABECALHO, de_link, de_xlsx  # noqa: E402

CHAVES = RAIZ / "dados" / "ficha" / "analise-chaves.tsv"
AMOSTRA = RAIZ / "src" / "testes" / "dados" / "ficha_analise_branca.xlsx"


def _celula(texto) -> str:
    """A coluna C às vezes é fórmula de texto (="'FICHA PESSOAL'!D64"), pra a aspa não sumir no Sheets."""
    texto = str(texto or "")
    achou = re.fullmatch(r'="(.*)"', texto)
    return achou.group(1).replace('""', '"') if achou else texto


def linhas_do_xlsx(caminho: str) -> list[tuple[str, str, str, str]]:
    """(chave, célula, parte, valor do cabeçalho) de cada linha, lidas do .xlsx com as fórmulas."""
    from openpyxl import load_workbook

    livro = load_workbook(caminho, read_only=True)
    try:
        linhas = []
        for chave, valor, celula, parte in livro[ABA].iter_rows(min_row=2, max_col=4, values_only=True):
            if chave:
                fixo = "" if str(valor or "").startswith("=") else str(valor or "")
                linhas.append((str(chave), _celula(celula), str(parte or ""), fixo))
        return linhas
    finally:
        livro.close()


def chaves(fonte: str) -> None:
    if fonte.lower().endswith(".xlsx"):
        linhas = linhas_do_xlsx(fonte)
        origem = Path(fonte).name
    else:
        ficha = asyncio.run(de_link(fonte))
        linhas = [(k, c.celula, c.parte, str(c.valor) if c.parte == "cabeçalho" else "") for k, c in ficha.items()]
        origem = "link da ficha oficial"
    repetidas = {k for k in (l[0] for l in linhas) if [l[0] for l in linhas].count(k) > 1}
    if repetidas:
        sys.exit(f"chave repetida na aba: {sorted(repetidas)[:5]}")
    formato = next((l[3] for l in linhas if l[0] == "formato.versao"), "?")
    with open(CHAVES, "w", encoding="utf-8", newline="\n") as saida:
        saida.write(f"# Chaves da aba ANALISE, geradas por ferramentas/ficha_analise.py em {date.today():%d/%m/%Y}\n")
        saida.write(f"# fonte: {origem} · formato.versao {formato} · {len(linhas)} chaves. Não edite à mão.\n")
        saida.write("chave\tparte\tcélula\n")
        for chave, celula, parte, _ in linhas:
            saida.write(f"{chave}\t{parte}\t{celula}\n")
    partes = {}
    for _, _, parte, _ in linhas:
        partes[parte] = partes.get(parte, 0) + 1
    print(f"{CHAVES.relative_to(RAIZ)}: {len(linhas)} chaves {partes}")


def amostra(ficha: str, nome: str = AMOSTRA.name, *preencher: str) -> None:
    """Recalcula a ficha no LibreOffice e grava só a ANALISE com valores, como sairia do download do Google."""
    from openpyxl import Workbook, load_workbook

    if not shutil.which("libreoffice"):
        sys.exit("precisa do LibreOffice pra recalcular")
    pasta = Path(tempfile.mkdtemp(prefix="analise-"))
    try:
        copia = pasta / "ficha.xlsx"
        shutil.copy(ficha, copia)
        livro = load_workbook(copia)
        for aba in livro:  # o LibreOffice só reconhece o IFS com o prefixo do Excel (como no ler-analise.py)
            for celula in tuple(aba._cells.values()):
                if isinstance(celula.value, str) and celula.value.startswith("=") and "IFS(" in celula.value \
                        and "_xlfn.IFS(" not in celula.value:
                    celula.value = re.sub(r"(?<![A-Z_.])IFS\(", "_xlfn.IFS(", celula.value)
        for item in preencher:  # "FICHA!D11=Bastião", "'FICHA PESSOAL'!AH64="
            alvo, _, valor = item.partition("=")
            aba_alvo, _, celula = alvo.rpartition("!")
            livro[aba_alvo.strip("'")][celula] = int(valor) if re.fullmatch(r"-?\d+", valor) else (valor or None)
        livro.save(copia)
        livro.close()
        saida = pasta / "saida"
        subprocess.run(["libreoffice", "--headless", "--convert-to", "xlsx", "--outdir", str(saida), str(copia)],
                       capture_output=True, timeout=900, check=False)
        feito = saida / "ficha.xlsx"
        if not feito.exists():
            sys.exit("o LibreOffice não converteu a ficha")
        recalculada = load_workbook(feito, data_only=True, read_only=True)
        linhas = [list(l) for l in recalculada[ABA].iter_rows(max_col=4, values_only=True)]
        recalculada.close()
    finally:
        shutil.rmtree(pasta, ignore_errors=True)
    novo = Workbook()
    aba = novo.active
    aba.title = ABA
    for linha in linhas:
        aba.append([None if v is None else str(v) for v in linha])  # tudo texto, como a ficha publica
    aba.sheet_state = "hidden"
    novo.create_sheet("FICHA")  # um livro não salva com todas as abas ocultas
    novo.move_sheet("FICHA", offset=-1)
    destino = AMOSTRA.parent / nome
    destino.parent.mkdir(parents=True, exist_ok=True)
    novo.save(destino)
    lida = de_xlsx(destino)
    print(f"{destino.relative_to(RAIZ)}: {len(lida)} chaves, {os.path.getsize(destino) // 1024} KB")


if __name__ == "__main__":
    if len(sys.argv) < 3 or sys.argv[1] not in ("chaves", "amostra") or (sys.argv[1] == "chaves" and len(sys.argv) != 3):
        sys.exit(__doc__)
    {"chaves": chaves, "amostra": amostra}[sys.argv[1]](*sys.argv[2:])
    if tuple(CABECALHO) != ("chave", "valor", "célula", "parte"):
        sys.exit("cabeçalho do leitor mudou")
