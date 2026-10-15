"""Prepara cinco arquivos sintéticos para testar o painel no Discord, sem API.

As fontes guardam só a ANALISE recalculada pelo gerador/LibreOffice, não uma ficha
editável completa. A saída vai para src/var/, fora do Git.
"""
from pathlib import Path
import shutil
import sys

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "src"))
from closure_bot.ficha_corretor import Referencias, conferir
from closure_bot.ficha_leitor import de_xlsx


def preparar():
    destino = RAIZ / "src/var/testes-tr-pericias"
    destino.mkdir(parents=True, exist_ok=True)
    fontes = RAIZ / "src/testes/dados"
    nomes = {"01-pessoal-verde.xlsx": ("pessoal", {}, "aprovada", set()),
             "02-especialista-corda-amarelo.xlsx": ("especialista", {}, "aprovada", {"item.1.volume", "carga.total"}),
             "03-domada-verde.xlsx": ("domada", {}, "aprovada", set()),
             "04-tr-adulterado-vermelho.xlsx": ("domada", {"invocacao.1.tr.vigor.bonus": 999}, "reprovada", set()),
             "05-tr-faltante-amarelo.xlsx": ("pessoal", {"tr.vigor.treino": 0, "tr.vigor.bonus": 2}, "aprovada", {"tr"})}
    ref = Referencias.do_livro(RAIZ / "dados")
    for nome, (caso, trocas, veredito, pendencias) in nomes.items():
        saida = destino / nome
        shutil.copyfile(fontes / f"ficha_analise_v4_{caso}.xlsx", saida)
        if trocas:
            w = load_workbook(saida)
            try:
                for linha in w["ANALISE"].iter_rows(min_row=2):
                    if linha[0].value in trocas:
                        linha[1].value = trocas[linha[0].value]
                w.save(saida)
            finally:
                w.close()
        r = conferir(de_xlsx(saida), ref)
        assert r.veredito == veredito, nome
        assert {p.campo for p in r.pendencias} == pendencias, nome
        if trocas.get("invocacao.1.tr.vigor.bonus"):
            assert {p.campo for p in r.problemas} == {"invocacao.1.tr.vigor.bonus"}, nome
        else:
            assert not r.problemas, nome
        print(nome + ": " + ("vermelho" if r.problemas else "amarelo" if r.pendencias else "verde"))
    return destino


if __name__ == "__main__":
    preparar()
