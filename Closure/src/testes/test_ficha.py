"""Fase 3: leitor da aba ANALISE e corretor da ficha.

As fichas de teste em testes/dados/ficha_analise_*.xlsx são a ficha gerada (ficha-projeto-m-0.1.xlsx, B47 fase 1),
preenchida e recalculada no LibreOffice pela ferramentas/ficha_analise.py: os resultados que o corretor confere foram
calculados pela própria ficha, não escritos aqui. A Kaori é a ficha de exemplo do livro (peça 8).

Precisa do openpyxl (sem ele, pula). O teste pelo link da ficha oficial só roda com FICHA_ONLINE=1.
"""
import asyncio
import io
import os
import tempfile
import unittest
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

from closure_bot.ficha_corretor import (LIVRO, Aprovados, Referencias, RegrasDaCampanha, classificar_aviso, conferir,
                                        impressao_digital, refino)
from closure_bot.ficha_leitor import Campo, FichaIlegivel, de_csv, de_link, de_xlsx, id_da_planilha, normaliza, resumo

DADOS = Path(__file__).parent / "dados"
RAIZ = Path(__file__).resolve().parents[2]
FICHA_OFICIAL = "https://docs.google.com/spreadsheets/d/1u0zWQFVaRTTc_VlrAJsabgAiJ1D4rTYhzG7QQWof1ro/edit"
FICHA_GERADA = Path("/media/mizuki/HD Externo II/Claude/Claude 3/.claude/worktrees/ficha-invocacao-update-a350cb/"
                    "ficha-v01/ficha-projeto-m-0.1.xlsx")


def _ficha(nome):
    return de_xlsx(DADOS / f"ficha_analise_{nome}.xlsx")


def _adulterada(nome, chave, valor):
    """Cópia da ficha com o valor de uma chave trocado à mão na aba ANALISE, como faria quem mexe na própria cópia."""
    livro = openpyxl.load_workbook(DADOS / f"ficha_analise_{nome}.xlsx")
    for linha in livro["ANALISE"].iter_rows(min_row=2, max_col=2):
        if linha[0].value == chave:
            linha[1].value = valor
    memoria = io.BytesIO()
    livro.save(memoria)
    return de_xlsx(memoria.getvalue())


def _contas(resultado):
    return {p.campo: (p.esperado, p.encontrado) for p in resultado.problemas if p.regra == "conta"}


def _com_campos(arquivo, **trocas):
    """Entradas normalizadas da ficha real, com as alterações do caso de regressão."""
    ficha = _ficha(arquivo)
    for chave, valor in trocas.items():
        chave = chave.replace("__", ".")
        anterior = ficha[chave]
        ficha[chave] = Campo(normaliza(valor), anterior.celula, anterior.parte)
    return ficha


class Normalizacao(unittest.TestCase):
    def test_valores_da_planilha(self):
        casos = {"1,5": 1.5, "3": 3, "-2": -2, "VERDADEIRO": 1, "FALSE": 0, None: "", "  Latente ": "Latente", "1.0": 1.0}
        for entrada, esperado in casos.items():
            self.assertEqual(normaliza(entrada), esperado, entrada)

    def test_csv_e_pagina_de_login(self):
        texto = '"chave","valor","célula","parte"\n"formato.versao","1","","cabeçalho"\n"defesa","11","FICHA!D59","resultado"\n'
        ficha = de_csv(texto)
        self.assertEqual((ficha["formato.versao"].valor, ficha["defesa"].valor, ficha["defesa"].celula), (1, 11, "FICHA!D59"))
        with self.assertRaises(FichaIlegivel):
            de_csv("<!DOCTYPE html><html>faça login</html>")
        with self.assertRaises(FichaIlegivel):
            de_csv('"a","b"\n"c","d"\n')

    def test_id_da_planilha(self):
        self.assertEqual(id_da_planilha(FICHA_OFICIAL), "1u0zWQFVaRTTc_VlrAJsabgAiJ1D4rTYhzG7QQWof1ro")
        for ruim in ("https://exemplo.com/planilha", "abc", "https://docs.google.com/spreadsheets/d/x/edit"):
            with self.assertRaises(FichaIlegivel):
                id_da_planilha(ruim)


@unittest.skipIf(openpyxl is None, "openpyxl não instalado")
class Leitor(unittest.TestCase):
    def test_le_a_ficha_gerada_com_as_2650_chaves(self):
        ficha = _ficha("branca")
        self.assertEqual(len(ficha), 2650)
        self.assertEqual(resumo(ficha), {"cabeçalho": 4, "entrada": 1629, "resultado": 601, "aviso": 416})
        self.assertEqual((ficha["formato.versao"].valor, ficha["ficha.livro"].valor), (1, "Ciclo Maldito R51"))
        for chave in ("sequela.1", "pericia.atletismo.treino", "tr.fisico.treino"):
            self.assertIn(ficha[chave].valor, (0, 1), chave)
        self.assertEqual(ficha["em_uso.vestindo"].celula, "'FICHA PESSOAL'!AH64")

    def test_contrato_tem_as_mesmas_chaves_da_ficha(self):
        linhas = (RAIZ / "dados" / "ficha" / "analise-chaves.tsv").read_text(encoding="utf-8").splitlines()
        chaves = [l.split("\t")[0] for l in linhas if l and not l.startswith("#") and not l.startswith("chave\t")]
        self.assertEqual(chaves, list(_ficha("branca")))

    def test_sem_aba_ou_sem_valores_nao_e_lida(self):
        sem_aba = openpyxl.Workbook()
        memoria = io.BytesIO()
        sem_aba.save(memoria)
        with self.assertRaisesRegex(FichaIlegivel, "ANALISE"):
            de_xlsx(memoria.getvalue())
        with self.assertRaises(FichaIlegivel):
            de_xlsx(b"isso nao e uma planilha")
        if FICHA_GERADA.exists():  # a ficha saída do gerador tem só as fórmulas, sem os resultados
            with self.assertRaisesRegex(FichaIlegivel, "valores calculados"):
                de_xlsx(FICHA_GERADA)


@unittest.skipIf(openpyxl is None, "openpyxl não instalado")
class Corretor(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.ref = Referencias.do_livro(RAIZ / "dados")

    def test_referencias_saem_do_livro(self):
        self.assertEqual(self.ref.origens, ("Latente", "Receptáculo", "Descendente", "Encarnado", "Híbrido",
                                            "Corpo Amaldiçoado", "Restrição Celestial"))
        self.assertEqual((self.ref.marcos, self.ref.refino_maximo), ((6, 10, 14, 18, 22, 26, 30), 10))
        self.assertEqual((self.ref.uniformes["Revestimento 3"].protecao, self.ref.uniformes["Revestimento 3"].teto_destreza), (6, 0))
        self.assertIsNone(self.ref.uniformes["Traje 2"].teto_destreza)
        self.assertEqual(self.ref.escudos["Torre"].teto_destreza, 1)
        self.assertEqual([refino(n, m, self.ref) for n, m in ((2, 0), (6, 0), (6, 1), (14, 2), (30, 7))], [1, 2, 3, 6, 10])

    def test_ficha_em_branco_sai_reprovada_pelas_pericias_e_pelo_caminho_por_escolher(self):
        resultado = conferir(_ficha("branca"), self.ref)
        self.assertEqual(resultado.veredito, "reprovada")
        self.assertEqual([(p.regra, p.campo) for p in resultado.problemas],
                         [("escolha", "caminho"), ("escolha", "trilha"), ("aviso", "aviso.escolhas_de_pericia")])
        self.assertTrue(resultado.problemas[-1].encontrado.startswith("Falta: 9 perícias"))
        self.assertEqual(resultado.conferidos, ["Origem Latente", "Integridade 25", "Proteção 1", "Defesa 11"])
        self.assertIn("vida e PE: sem Caminho do livro escolhido", resultado.nao_conferido)

    def test_kaori_bate_com_o_livro_e_com_a_ficha(self):
        ficha = _ficha("kaori")
        self.assertEqual({k: ficha[k].valor for k in ("vida_max", "integridade_max", "energia_max", "defesa")},
                         {"vida_max": 23, "integridade_max": 26, "energia_max": 8, "defesa": 13})  # peça 8, "Os números"
        resultado = conferir(ficha, self.ref)
        self.assertEqual(_contas(resultado), {})
        self.assertEqual([p.regra for p in resultado.problemas if p.regra in ("nome", "escolha")], [])
        self.assertEqual(resultado.conferidos, ["Caminho Bastião", "Trilha Muro", "Origem Descendente", "Vida 23", "PE 8",
                                                "Integridade 26", "Proteção 1", "Defesa 13"])

    def test_buff_entra_na_conta_e_vai_pro_mestre_conferir(self):
        notas = [(p.regra, p.campo, p.encontrado) for p in conferir(_ficha("revestimento"), self.ref).notas if p.regra == "buff"]
        self.assertEqual(notas, [("buff", "buff.defesa", "1"), ("buff", "buff.vida", "2")])

    def test_uniforme_escudo_marcos_e_buff_batem_com_a_ficha(self):
        for nome in ("revestimento", "refino"):  # Revestimento 2 + Broquel + Buff; Refino 6 pelos marcos + escudo Médio
            self.assertEqual(_contas(conferir(_ficha(nome), self.ref)), {}, nome)

    def test_defesa_trocada_a_mao_acende_o_problema_certo(self):
        resultado = conferir(_adulterada("branca", "defesa", "15"), self.ref)
        self.assertEqual(_contas(resultado), {"defesa": ("11", "15")})
        problema = next(p for p in resultado.problemas if p.campo == "defesa")
        self.assertEqual((problema.regra, problema.ancora), ("conta", "equip--eqp-protecao"))

    def test_vida_e_pe_trocados_a_mao_na_kaori(self):
        resultado = conferir(_adulterada("kaori", "vida_max", "30"), self.ref)
        self.assertEqual(_contas(resultado), {"vida_max": ("23", "30")})
        self.assertEqual(_contas(conferir(_adulterada("kaori", "energia_max", ""), self.ref)), {"energia_max": ("8", "vazio")})

    def test_vida_rolada_na_campanha_nao_e_conferida(self):
        rolada = RegrasDaCampanha(nome="mesa do Sui", vida_por_nivel="rolada")
        resultado = conferir(_adulterada("kaori", "vida_max", "30"), self.ref, rolada)
        self.assertNotIn("vida_max", _contas(resultado))
        self.assertIn("vida: a campanha rola a vida por nível", resultado.nao_conferido)
        self.assertEqual(LIVRO.vida_por_nivel, "fixa")

    def test_formato_desconhecido_avisa_sem_corrigir(self):
        resultado = conferir(_adulterada("kaori", "formato.versao", "2"), self.ref)
        self.assertEqual(resultado.veredito, "formato desconhecido")
        self.assertEqual([(p.regra, p.encontrado) for p in resultado.problemas], [("formato", "2")])

    def test_nomes_contra_o_catalogo_e_o_capitulo_de_origens(self):
        casos = {
            ("origem", "Feto"): "origem", ("origem", "Corpo Amaldiçoado · Sem Técnica"): "origem",
            ("origem", "Restrição Celestial"): "origem", ("trilha", "Estocada"): "trilha", ("caminho", "Paladino"): "caminho",
        }
        for (chave, valor), campo in casos.items():
            nomes = [p.campo for p in conferir(_adulterada("kaori", chave, valor), self.ref).problemas if p.regra == "nome"]
            self.assertEqual(nomes, [campo], valor)
        sem_origem = [p for p in conferir(_adulterada("kaori", "origem", ""), self.ref).problemas if p.regra == "escolha"]
        self.assertEqual([(p.campo, p.esperado) for p in sem_origem], [("origem", "escolher uma Origem")])
        for valor in ("Latente · Sem Técnica", "Restrição Celestial · Sem Energia", "Híbrido"):
            nomes = [p for p in conferir(_adulterada("kaori", "origem", valor), self.ref).problemas if p.regra == "nome"]
            self.assertEqual(nomes, [], valor)

    def _com(self, nome, **trocas):
        """Ficha de teste com várias chaves trocadas à mão."""
        livro = openpyxl.load_workbook(DADOS / f"ficha_analise_{nome}.xlsx")
        for linha in livro["ANALISE"].iter_rows(min_row=2, max_col=2):
            chave = linha[0].value.replace(".", "__")
            if chave in trocas:
                linha[1].value = trocas[chave]
        memoria = io.BytesIO()
        livro.save(memoria)
        return de_xlsx(memoria.getvalue())

    def test_treino_extra_da_origem_nao_reprova_e_vai_pro_mestre(self):
        # Encarnado com Duas Vidas: um treino a mais pelo Conhecimento Reavivado, que a ficha conta como "1 a mais"
        ficha = self._com("kaori", origem="Encarnado", origem__marcante="Duas Vidas",
                          aviso__escolhas_de_pericia="Marcou 1 a mais: tire uma perícia, um ofício ou uma especialização",
                          aviso__pericias_disponiveis="1 a mais")
        resultado = conferir(ficha, self.ref)
        self.assertEqual([p.campo for p in resultado.problemas if p.regra == "aviso"], [])
        self.assertEqual({p.campo for p in resultado.notas if p.regra == "origem"},
                         {"aviso.escolhas_de_pericia"})
        dois = self._com("kaori", origem="Encarnado", origem__marcante="Duas Vidas", aviso__pericias_disponiveis="2 a mais")
        self.assertIn("aviso.pericias_disponiveis", [p.campo for p in conferir(dois, self.ref).problemas])
        # Especialidade Antiga: a especialização grátis só a partir do nível 10
        nivel9 = self._com("kaori", nivel="9", origem__menor_1="Especialidade Antiga", aviso__pericias_disponiveis="1 a mais")
        self.assertIn("aviso.pericias_disponiveis", [p.campo for p in conferir(nivel9, self.ref).problemas])
        nivel10 = self._com("kaori", nivel="10", origem__menor_1="Especialidade Antiga", aviso__pericias_disponiveis="1 a mais")
        self.assertNotIn("aviso.pericias_disponiveis", [p.campo for p in conferir(nivel10, self.ref).problemas])

    def test_pe_da_origem_que_a_ficha_nao_soma_pede_o_buff(self):
        # Núcleo Dependente, nível 2 e Essência 1: +1 PE (Essência, sem marco)
        sem_buff = conferir(self._com("kaori", origem__menor_1="Núcleo Dependente"), self.ref)
        self.assertEqual([(p.campo, p.esperado) for p in sem_buff.pendencias if p.regra == "origem"],
                         [("energia_max", "+1 de Núcleo Dependente no Buff de PE")])
        com_buff = conferir(self._com("kaori", origem__menor_1="Núcleo Dependente", buff__energia="1", energia_max="9"), self.ref)
        self.assertEqual([p for p in com_buff.pendencias if p.regra == "origem"], [])
        self.assertEqual(_contas(com_buff), {})

    def test_escudo_nao_conta_com_arma_de_duas_maos(self):
        # Incursor com escudo Médio: proteção 5 e Defesa 18; com uma Besta na mão principal, o escudo sai: 3 e 16
        com_besta = self._com("refino", em_uso__principal="Besta", protecao="3", defesa="16")
        self.assertEqual(_contas(conferir(com_besta, self.ref)), {})
        self.assertEqual(_contas(conferir(self._com("refino", protecao="3", defesa="16"), self.ref)),
                         {"protecao": ("5", "3"), "defesa": ("18", "16")})
        self.assertIn("Kanabō", self.ref.duas_maos)

    def test_rota_sem_energia_confere_a_defesa_pela_lapidacao(self):
        # Defesa sem Armadura: 1 + Lapidação/3, a mesma conta e progressão do Refino; Kaori sem uniforme fica com 13
        resultado = conferir(_adulterada("kaori", "origem", "Restrição Celestial · Sem Energia"), self.ref)
        self.assertEqual(_contas(resultado), {})
        self.assertIn("Defesa 13", resultado.conferidos)

    def test_restricao_corporal_defesa_zero_e_vida_pela_metade(self):
        debilitado = conferir(self._com("kaori", origem="Restrição Celestial · Restrição corporal",
                                        origem__marcante="Corpo Debilitado"), self.ref)
        self.assertEqual([(p.campo, p.encontrado) for p in debilitado.pendencias if p.regra == "origem"], [("defesa", "13")])
        self.assertNotIn("defesa", _contas(debilitado))
        ajustada = conferir(self._com("kaori", origem="Restrição Celestial · Restrição corporal", origem__marcante="Corpo Debilitado",
                                      defesa="0"), self.ref)
        self.assertIn("Defesa 0 (Corpo Debilitado)", ajustada.conferidos)
        quebradico = conferir(self._com("kaori", origem="Restrição Celestial · Restrição corporal",
                                        origem__marcante="Corpo Quebradiço", vida_max="11", defesa="0"), self.ref)
        self.assertIn("Vida 11 (metade, Corpo Quebradiço)", quebradico.conferidos)  # 23 // 2
        self.assertEqual(_contas(quebradico), {})

    def test_defesa_zero_distingue_formula_normal_de_adulteracao(self):
        # Reprodução local do Teste 3: Incursor 14, Besta, escudo Médio sem efeito, Defesa normal 15.
        for marcante in ("Corpo Debilitado", "Corpo Quebradiço", "Corpo Exaurido"):
            for valor in (0, 15, 19, ""):
                with self.subTest(marcante=marcante, defesa=valor):
                    ficha = _com_campos("refino", origem="Restrição Celestial · Restrição corporal",
                                        origem__marcante=marcante, em_uso__principal="Besta",
                                        atributo__destreza__base=2, atributo__destreza__corpo=0,
                                        protecao=3, defesa=valor)
                    resultado = conferir(ficha, self.ref)
                    pendencias = [p for p in resultado.pendencias if p.campo == "defesa"]
                    if valor == 0:
                        self.assertNotIn("defesa", _contas(resultado))
                        self.assertEqual(pendencias, [])
                        self.assertIn(f"Defesa 0 ({marcante})", resultado.conferidos)
                    elif valor == 15:
                        self.assertNotIn("defesa", _contas(resultado))
                        self.assertEqual([p.encontrado for p in pendencias], ["15"])
                    else:
                        self.assertEqual(resultado.veredito, "reprovada")
                        self.assertEqual(_contas(resultado)["defesa"], ("0", str(valor) or "vazio"))
                        self.assertEqual(pendencias, [])

    def test_defesa_zero_continua_conferindo_protecao(self):
        ficha = _com_campos("kaori", origem="Restrição Celestial · Restrição corporal",
                            origem__marcante="Corpo Debilitado", defesa=0, protecao=9)
        resultado = conferir(ficha, self.ref)
        self.assertEqual(_contas(resultado), {"protecao": ("1", "9")})
        self.assertEqual(resultado.veredito, "reprovada")

    def test_buff_nao_muda_defesa_fixada_e_formula_conhecida_ainda_avisa(self):
        for valor in (0, 17, 19):
            ficha = _com_campos("kaori", origem="Restrição Celestial · Restrição corporal",
                                origem__marcante="Corpo Debilitado", buff__defesa=4, defesa=valor)
            resultado = conferir(ficha, self.ref)
            self.assertEqual([p.encontrado for p in resultado.notas if p.campo == "buff.defesa"], ["4"])
            if valor == 0:
                self.assertIn("Defesa 0 (Corpo Debilitado)", resultado.conferidos)
            elif valor == 17:
                self.assertNotIn("defesa", _contas(resultado))
                self.assertEqual([p.encontrado for p in resultado.pendencias if p.campo == "defesa"], ["17"])
            else:
                self.assertEqual(_contas(resultado)["defesa"], ("0", "19"))

    def test_tres_avisos_do_mesmo_treino_viram_uma_nota(self):
        ficha = _com_campos("kaori", origem="Encarnado", origem__marcante="Duas Vidas",
                            aviso__pericias_disponiveis="1 a mais", aviso__oficios_disponiveis="1 a mais",
                            aviso__escolhas_de_pericia="Marcou 1 a mais: tire uma perícia, um ofício ou uma especialização")
        for ordenada in (ficha, dict(reversed(list(ficha.items())))):
            resultado = conferir(ordenada, self.ref)
            self.assertEqual([p.campo for p in resultado.notas if p.regra == "origem"], ["aviso.escolhas_de_pericia"])
            self.assertEqual([p for p in resultado.problemas if p.regra == "aviso"], [])

    def test_agrupamento_nao_apaga_excesso_sem_permissao_ou_contadores_diferentes(self):
        ficha = _com_campos("kaori", origem="Encarnado", origem__marcante="Duas Vidas",
                            aviso__pericias_disponiveis="2 a mais", aviso__oficios_disponiveis="1 a mais",
                            aviso__escolhas_de_pericia="Marcou 1 a mais")
        resultado = conferir(ficha, self.ref)
        self.assertEqual([p.campo for p in resultado.problemas if p.regra == "aviso"], ["aviso.pericias_disponiveis"])
        self.assertEqual([p.campo for p in resultado.notas if p.regra == "origem"], ["aviso.escolhas_de_pericia"])
        ficha = _com_campos("refino", origem="Encarnado", origem__marcante="Duas Vidas",
                            origem__menor_1="Especialidade Antiga",
                            aviso__pericias_disponiveis="2 a mais", aviso__escolhas_de_pericia="Marcou 1 a mais")
        resultado = conferir(ficha, self.ref)
        self.assertEqual({p.campo for p in resultado.notas if p.regra == "origem"},
                         {"aviso.pericias_disponiveis", "aviso.escolhas_de_pericia"})


class Avisos(unittest.TestCase):
    def test_classificacao(self):
        casos = {
            ("aviso.escolhas_de_pericia", "Falta: 9 perícias e 2 ofícios, ou 10 perícias sem ofício"): "erro",
            ("aviso.pontos_disponiveis", "3 a mais"): "erro", ("aviso.pontos_disponiveis", "Um atributo passou de 6"): "erro",
            ("aviso.pontos_disponiveis", "Na criação, nenhum acima de 3"): "erro", ("aviso.grau", "Pede Grau 3"): "erro",
            ("feitico.3.erros", "faltam 2 pontos"): "erro", ("invocacao.1.pontos_acima", "2"): "erro",
            ("aviso.pontos_disponiveis", "2 de 11"): "pendente", ("aviso.pontos_disponiveis", "0 de 9"): None,
            ("aviso.testes_disponiveis", "2 de 2"): "pendente", ("aviso.pericias_disponiveis", "9 de 9"): None,
            ("aviso.forca_v", "Falta Força: 0 de 4"): "pendente", ("aviso.requisito", "Não cumprido"): "pendente",
            ("aviso.treino", "Sem treino: desvantagem"): "pendente", ("aviso.selo", "Trava Selo de gesto"): "pendente",
            ("aviso.pontos_disponiveis", "Especialização só no nível 10"): "erro",
            ("feitico.3.avisos", "1 ponto sem Melhoria para pagar"): "nota",
            ("aviso.treino", "Treinada"): None, ("aviso.grau", "Liberado no Grau 4"): None, ("aviso.selo", "Sem escudo"): None,
            ("aviso.qualquer", "texto que a ficha ainda não tinha"): "nota", ("aviso.forca_s", "—"): None,
        }
        for (chave, texto), esperado in casos.items():
            self.assertEqual(classificar_aviso(chave, texto), esperado, (chave, texto))


class ConteudoCustom(unittest.TestCase):
    def test_impressao_digital_ignora_forma_e_pega_mudanca_de_texto(self):
        campos = {"dano": "1d8 Cortante", "propriedades": "Fineza"}
        base = impressao_digital("arma", "Lâmina da Avó", campos)
        self.assertEqual(base, impressao_digital("Arma", "  lamina   da avo ", {"propriedades": "fineza", "dano": "1D8 cortante"}))
        self.assertNotEqual(base, impressao_digital("arma", "Lâmina da Avó", {**campos, "dano": "1d10 Cortante"}))
        aprovados = Aprovados()
        aprovados.aprovar("arma", "Lâmina da Avó", campos, por="mestre 1, 14/10/2026")
        self.assertTrue(aprovados.aprovado("arma", "lâmina da avó", campos))
        self.assertFalse(aprovados.aprovado("arma", "Lâmina da Avó", {**campos, "propriedades": "Fineza, Arremesso"}))


class Comando(unittest.IsolatedAsyncioTestCase):
    @unittest.skipIf(openpyxl is None, "openpyxl não instalado")
    def test_painel_sem_nome_usa_caminho_e_nivel_so_quando_conhecidos(self):
        from types import SimpleNamespace
        from closure_bot.cogs.ficha import painel
        ref = Referencias.do_livro(RAIZ / "dados")
        persona = SimpleNamespace(chave="closure", nome="Closure", avatar_url=None)
        ficha = _com_campos("kaori", nome="")
        self.assertEqual(painel(conferir(ficha, ref), ficha, persona).title, "Ficha de Bastião, nível 2")
        ficha = _com_campos("kaori", nome="", caminho="Escolha seu Caminho")
        self.assertEqual(painel(conferir(ficha, ref), ficha, persona).title, "Ficha sem nome")
        ficha = _com_campos("kaori", nome="", nivel="")
        self.assertEqual(painel(conferir(ficha, ref), ficha, persona).title, "Ficha de Bastião")

    async def test_ligado_registra_e_desligado_nao(self):
        from unittest.mock import AsyncMock, MagicMock, patch
        from closure_bot.cogs import ficha as modulo
        self.assertTrue(modulo.LIGADO)
        bot = MagicMock()
        bot.add_cog = AsyncMock()
        bot.config.dados = RAIZ / "dados"
        with patch.object(modulo, "LIGADO", False):
            await modulo.setup(bot)
        bot.add_cog.assert_not_called()
        await modulo.setup(bot)
        cog = bot.add_cog.call_args.args[0]
        self.assertEqual([c.name for c in cog.grupo.commands], ["analisar"])
        # sem o livro no pacote, a correção não sobe e o resto do bot segue
        quebrado = MagicMock()
        quebrado.add_cog = AsyncMock()
        quebrado.config.dados = Path(tempfile.mkdtemp())
        with self.assertLogs("closure_bot.cogs.ficha", level="ERROR"):
            await modulo.setup(quebrado)
        quebrado.add_cog.assert_not_called()

    @unittest.skipIf(openpyxl is None, "openpyxl não instalado")
    def test_painel_diz_o_que_falta_o_que_bateu_e_quem_fala(self):
        from types import SimpleNamespace
        from closure_bot.cogs.ficha import FALA, painel
        ref = Referencias.do_livro(RAIZ / "dados")
        ficha = _adulterada("kaori", "defesa", "15")
        closure = SimpleNamespace(chave="closure", nome="Closure", avatar_url=None)
        embed = painel(conferir(ficha, ref), ficha, closure, {"edicao_livro": "R51"})
        campos = {c.name: c.value for c in embed.fields}
        self.assertEqual((embed.title, embed.author.name, embed.description), ("Ficha de Kaori", "Closure", FALA["closure"]["reprovada"]))
        self.assertIn("**Defesa**: está 15, o certo é 13", campos["❌ Precisa corrigir"])
        self.assertIn("Vida 23", campos["✅ Conferido e certo"])
        self.assertIn("feitiços e invocações", campos["⏳ Ainda não conferido"])
        self.assertIn("R51", embed.footer.text)
        cheia = painel(conferir(_ficha("revestimento"), ref), _ficha("revestimento"), closure)
        campos_cheia = {c.name: c.value for c in cheia.fields}
        self.assertIn("Buff/Debuff de defesa", campos_cheia["👀 Pro mestre conferir"])
        falta = campos_cheia["⚠️ Falta ou pede atenção"]
        for trecho in ("**Pontos de atributo**: faltam distribuir 2 de 11", "**Testes de Resistência**: faltam escolher 2 de 2",
                       "**Força do uniforme**: Falta Força: 0 de 4"):
            self.assertIn(trecho, falta)
        self.assertTrue(all(len(c.value) <= 1024 for c in cheia.fields))


@unittest.skipUnless(os.environ.get("FICHA_ONLINE") == "1", "precisa de internet: rode com FICHA_ONLINE=1")
class PeloLink(unittest.TestCase):
    def test_ficha_oficial_pelo_link_tem_as_2650_chaves(self):
        ficha = asyncio.run(de_link(FICHA_OFICIAL))
        self.assertEqual(len(ficha), 2650)
        self.assertEqual(ficha["formato.versao"].valor, 1)
        self.assertEqual(conferir(ficha, Referencias.do_livro(RAIZ / "dados")).veredito, "reprovada")


if __name__ == "__main__":
    unittest.main()
