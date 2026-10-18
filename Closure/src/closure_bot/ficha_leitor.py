"""Leitor da aba oculta ANALISE da ficha do Ciclo Maldito (fase 3, contrato em dados/ficha/CONTRATO-ABA-BOT.md).

A aba tem quatro colunas, chave · valor · célula · parte, e a linha 1 é o cabeçalho. Duas entradas:
  - o .xlsx que o jogador baixa (Arquivo > Fazer download > .xlsx): o Google grava o resultado das fórmulas, e o
    openpyxl lê com data_only, mesmo com a aba oculta;
  - o link da planilha aberta pra leitura: a aba sai em CSV pela exportação do Google (gviz), sem conta e sem chave.

A coluna de valores vem toda como texto (no CSV o Google escolhe um tipo por coluna, e misturar número com texto
apagava os textos), então `normaliza` converte: número com vírgula decimal, caixa de marcar como 1/0 e vazio como
"não preenchido" (""). Referência: o ler-analise.py da worktree da ficha (B47, fase 1).

Este módulo só lê. Quem confere é o ficha_corretor.py, sem IA, e o bot nunca escreve na ficha.
"""
from __future__ import annotations

import csv
import io
import re
import math
import asyncio
from xml.etree import ElementTree as ET

from .ficha_esquema import tipo
from dataclasses import dataclass
from pathlib import Path

ABA = "ANALISE"
CABECALHO = ("chave", "valor", "célula", "parte")
# A aba é pública na ficha aberta pra leitura; o id da planilha é o que vem depois de /d/ no link.
ID_PLANILHA = re.compile(r"^[A-Za-z0-9_-]{20,}$")
URL_CSV = "https://docs.google.com/spreadsheets/d/{id}/gviz/tq?tqx=out:csv&sheet=" + ABA
LIMITE_BYTES = 5_000_000  # a ANALISE inteira tem uns 300 KB; mais que isso não é a aba


class FichaIlegivel(Exception):
    """A ficha não deu pra ler: sem a aba, sem valores, link fechado ou fora do formato. A mensagem vai pra pessoa."""


@dataclass(frozen=True)
class Campo:
    valor: object  # int, float ou str; "" é vazio
    celula: str
    parte: str


def normaliza(valor, chave: str | None = None) -> object:
    """'1,5' -> 1.5, '3' -> 3, 'VERDADEIRO'/'TRUE' -> 1, 'FALSO'/'FALSE' -> 0, None -> "", o resto texto sem as pontas."""
    if chave is not None and tipo(chave) == "texto":
        return "" if valor is None else str(valor)
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return valor if chave is not None and tipo(chave) == 'numero' else int(valor)
    if isinstance(valor, int):
        return valor
    if isinstance(valor, float):
        return int(valor) if float(valor).is_integer() else float(valor)
    texto = str(valor).strip()
    if texto.upper() in ("VERDADEIRO", "TRUE") and (chave is None or tipo(chave) == 'booleano'):
        return 1
    if texto.upper() in ("FALSO", "FALSE") and (chave is None or tipo(chave) == 'booleano'):
        return 0
    if re.fullmatch(r"-?\d+", texto):
        return int(texto)
    if re.fullmatch(r"-?\d+[,.]\d+", texto):
        return float(texto.replace(",", "."))
    return texto


def de_linhas(linhas) -> dict[str, Campo]:
    """[[chave, valor, célula, parte], ...], com ou sem o cabeçalho -> {chave: Campo}."""
    ficha: dict[str, Campo] = {}
    erros = {"#REF!", "#N/A", "#N/D", "#NAME?", "#NOME?", "#VALUE!", "#VALOR!", "#DIV/0!", "#NUM!", "#NULL!", "#ERROR!", "#SPILL!", "#CALC!"}
    erros_encontrados = []
    for n, linha in enumerate(linhas):
        if n > 20000:
            raise FichaIlegivel("a aba ANALISE excede o limite de linhas")
        linha = (list(linha) + ["", "", "", ""])[:4]
        chave = str(linha[0] or "").strip()
        if not chave or chave == "chave":
            continue
        if chave in ficha:
            raise FichaIlegivel(f"chave duplicada na ANALISE: {chave}")
        if str(linha[1]).strip().upper() in erros:
            erros_encontrados.append(f"{chave}: {linha[1]}")
        try:
            valor = normaliza(linha[1], chave)
            if isinstance(valor, (int, float)) and not math.isfinite(valor):
                raise ValueError('número fora do intervalo')
        except (ValueError, OverflowError) as erro:
            raise FichaIlegivel(f"número inválido ou fora do intervalo em {chave}") from erro
        ficha[chave] = Campo(valor, str(linha[2] or ""), str(linha[3] or ""))
    if erros_encontrados and ficha.get("formato.versao", Campo("", "", "")).valor in ("", 11):
        raise FichaIlegivel("erro de fórmula em " + erros_encontrados[0])
    return ficha


def de_csv(texto: str) -> dict[str, Campo]:
    """O CSV que o Google devolve pelo link. Página de login ou de erro no lugar do CSV vira FichaIlegivel."""
    if texto.lstrip()[:15].lower().startswith(("<!doctype", "<html")):
        raise FichaIlegivel("a planilha não está aberta pra leitura pelo link")
    ficha = de_linhas(csv.reader(io.StringIO(texto)))
    if "formato.versao" not in ficha:
        raise FichaIlegivel("a planilha não tem a aba ANALISE")
    return ficha


def de_xlsx(origem: str | Path | io.BytesIO | bytes) -> dict[str, Campo]:
    """O .xlsx baixado do Google. Ficha gerada e nunca aberta numa planilha não tem os valores das fórmulas."""
    from openpyxl import load_workbook  # só quem lê .xlsx precisa dele

    if isinstance(origem, (bytes, bytearray)):
        origem = io.BytesIO(origem)
    try:
        livro = load_workbook(origem, data_only=True, read_only=True)
    except Exception as erro:  # arquivo corrompido ou que não é .xlsx
        raise FichaIlegivel("o arquivo não é uma planilha .xlsx que dê pra abrir") from erro
    try:
        if ABA not in livro.sheetnames:
            raise FichaIlegivel("a planilha não tem a aba ANALISE (a ficha é de uma versão anterior?)")
        aba = livro[ABA]
        info = livro._archive.getinfo(aba._worksheet_path)
        if info.file_size > 20_000_000:
            raise FichaIlegivel("a aba ANALISE descompactada é grande demais")
        xml = ET.fromstring(livro._archive.read(aba._worksheet_path))
        ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        for c in xml.findall(".//s:sheetData/s:row/s:c", ns):
            if c.find("s:f", ns) is not None:
                v = c.find("s:v", ns)
                if (v is None or v.text is None) and c.get("t") != "str":
                    raise FichaIlegivel("o arquivo tem fórmula sem resultado calculado em ANALISE!" + c.get("r", ""))
        ficha = de_linhas(aba.iter_rows(max_col=4, values_only=True))
    finally:
        livro.close()
    if "formato.versao" not in ficha:
        raise FichaIlegivel("a aba ANALISE está fora do formato (falta formato.versao)")
    # O cabeçalho tem valores escritos; o resto é fórmula. Tudo vazio fora dele = o arquivo nunca foi calculado.
    if not any(c.valor != "" for c in ficha.values() if c.parte != "cabeçalho"):
        raise FichaIlegivel("o arquivo não tem os valores calculados: baixe a ficha pelo Google (Arquivo > Fazer download > .xlsx)")
    return ficha


def id_da_planilha(link: str) -> str:
    """O id a partir do link do Google Planilhas ou do próprio id. Outro endereço vira FichaIlegivel."""
    achou = re.search(r"docs\.google\.com/spreadsheets/d/([A-Za-z0-9_-]+)", link)
    pid = achou.group(1) if achou else link.strip()
    if not ID_PLANILHA.fullmatch(pid):
        raise FichaIlegivel("isso não parece um link do Google Planilhas")
    return pid


async def de_link(link: str, *, sessao=None, espera: float = 60) -> dict[str, Campo]:
    """Baixa a aba ANALISE pelo link aberto. `sessao` é uma aiohttp.ClientSession (o bot passa a dele)."""
    import aiohttp

    url = URL_CSV.format(id=id_da_planilha(link))
    propria = sessao is None
    sessao = sessao or aiohttp.ClientSession(timeout=aiohttp.ClientTimeout(total=espera))
    try:
        async with sessao.get(url) as resposta:
            if resposta.status != 200:
                raise FichaIlegivel("a planilha não está aberta pra leitura pelo link"
                                    if resposta.status in (401, 403, 404) else f"o Google respondeu {resposta.status}")
            corpo = b""
            # read(n) devolve o que já chegou, não o arquivo inteiro: lê em pedaços até o fim ou até o limite
            async for pedaco in resposta.content.iter_chunked(65536):
                corpo += pedaco
                if len(corpo) > LIMITE_BYTES:
                    break
    except (aiohttp.ClientError, asyncio.TimeoutError) as erro:
        raise FichaIlegivel("não deu pra baixar a planilha agora") from erro
    finally:
        if propria:
            await sessao.close()
    if len(corpo) > LIMITE_BYTES:
        raise FichaIlegivel("a resposta do link é grande demais pra ser a aba ANALISE")
    return de_csv(corpo.decode("utf-8", errors="replace"))


def resumo(ficha: dict[str, Campo]) -> dict[str, int]:
    """Quantas chaves de cada parte: {"entrada": 1629, ...}."""
    partes: dict[str, int] = {}
    for campo in ficha.values():
        partes[campo.parte] = partes.get(campo.parte, 0) + 1
    return partes
